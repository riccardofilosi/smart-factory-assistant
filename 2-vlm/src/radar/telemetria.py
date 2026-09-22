"""telemetria.py — strato osservatore della run. Il motore non viene toccato: qui
si usano solo primitive in sola lettura (board_rect_mask, build_model,
marker_angle, map_health, positions) e i dizionari già esposti da analyze/valuta.

Ogni metodo pubblico è avvolto in try/except: un errore di telemetria si logga e
la run non muore mai.

Artefatti per run (dentro la cartella-run):
  telemetria/00_mappa/        fasi costruzione mappa (PNG + mappa.json)
  telemetria/01_scarti/       scatti.csv (TUTTI gli scatti: esito+valori) + esempi foto
  telemetria/02_valutazioni/  vNNN_{blob,ncc_A,ncc_D,cosa}.png + vNNN_rec.json
  telemetria/gt_live.json     tasti 1/2/3 durante la run
  RUN_<runid>.xlsx            il deliverable: una riga per valutazione, miniature+link
"""
import csv
import datetime
import json
import os
import threading

import cv2
import numpy as np

# palette del disegno tecnico (4-relazione), in BGR
INK = (48, 36, 26)          # #1a2430
INK2 = (102, 89, 76)        # #4c5966
HAIR = (214, 208, 201)      # #c9d0d6
PAPER = (254, 253, 252)     # #fcfdfe
ACC = (10, 78, 191)         # #bf4e0a  caldo
ACC_D = (6, 51, 125)        # #7d3306
ACC_L = (187, 212, 242)     # #f2d4bb  debole
FONT = cv2.FONT_HERSHEY_SIMPLEX
MAX_ESEMPI = 3              # foto salvate per OGNI motivo di scarto (i numeri sempre)
THUMB_W = 150               # larghezza miniature nell'Excel


def _log(msg):
    print(f"[telemetria] {msg}", flush=True)


def _fantasma_uv(rec, df):
    """Fori-fantasma del record (posizione del KO precedente, ignorati dal giudizio)
    come set di uv, per disegnarli diversi dai pin scelti davvero."""
    out = set()
    for n in (rec.get("radar") or {}).get("fantasma") or []:
        try:
            out.add(df.parse_hole(n))
        except Exception:
            pass
    return out


def _rimozione_uv(rec, posx, df, att_uv):
    """Fori della zona del KO precedente (rec['rimozione'], solo frame di correzione):
    nel diff si accendono per la RIMOZIONE del pezzo, non perché il pezzo ci sia.
    Le viste li disegnano grigi. Ritorna set di uv, golden correnti esclusi."""
    rz = rec.get("rimozione") or {}
    out = set()
    for n in rz.get("fori") or []:
        try:
            out.add(df.parse_hole(n))
        except Exception:
            pass
    bb = rz.get("bbox")
    if bb:
        x0, y0, x1, y1 = bb
        g = 8
        for uv, p in posx.items():
            if x0 - g <= p[0] <= x1 + g and y0 - g <= p[1] <= y1 + g:
                out.add(uv)
    return out - set(att_uv)


def _blindato(fn):
    """Decoratore: la telemetria osserva, non comanda. Un suo errore non ferma la run.
    BaseException e non Exception: build_model su immagine storta fa sys.exit.
    Ctrl-C passa."""
    def wrap(*a, **k):
        try:
            return fn(*a, **k)
        except KeyboardInterrupt:
            raise
        except BaseException as e:
            _log(f"{fn.__name__} fallita ({type(e).__name__}: {str(e)[:80]}) - run prosegue")
            return None
    return wrap


class Telemetria:
    def __init__(self, rundir, cid=None, runid=None, sub="telemetria"):
        # `sub`: sottocartella degli artefatti. Serve a replay_vlm.py per rigiocare
        # una run senza sovrascrivere la telemetria live.
        self.rundir = os.path.abspath(rundir)
        self.cid = cid
        self.runid = runid or os.path.basename(self.rundir)
        self.base = os.path.join(self.rundir, sub)
        self.d_mappa = os.path.join(self.base, "00_mappa")
        self.d_scarti = os.path.join(self.base, "01_scarti")
        self.d_val = os.path.join(self.base, "02_valutazioni")
        for d in (self.d_mappa, self.d_scarti, self.d_val):
            os.makedirs(d, exist_ok=True)
        self.csv_path = os.path.join(self.d_scarti, "scatti.csv")
        if not os.path.exists(self.csv_path):
            with open(self.csv_path, "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerow(
                    ["nshot", "ora", "esito", "motivo", "d_prev", "novelty",
                     "allineamento", "fori_caldi", "blob_area", "foto"])
        self.n_esempi = {}          # motivo -> quante foto esempio già salvate
        self.contatori = {}         # esito/motivo -> conteggio (per il foglio run)
        self.righe = []             # una per valutazione -> foglio Excel
        self.gt = {}                # ordine -> giusto | errore_vero | falso_allarme
        self.mappa_info = {}
        self._gt_path = os.path.join(self.base, "gt_live.json")
        # `scatto` viene chiamata da due thread (il catturatore per gli scarti del
        # gate, il principale per i frame analizzati): il lock protegge la riga del
        # CSV e i contatori. `valutazione` e `chiudi` restano del solo thread principale.
        self._lock = threading.Lock()

    # ------------------------------------------------------------ 00: mappa
    @_blindato
    def fase_mappa(self, img0):
        """Rigioca la costruzione della mappa sullo scatto 0, una fase per PNG:
        1 maschera board -> 2 marker+griglia -> 3 eventuale deskew -> 4 salute.
        Stesse primitive di radar_live._mappa_analisi, in sola lettura."""
        import differential as df
        import build_map as bm
        import overlay_grid as og
        cv2.imwrite(os.path.join(self.d_mappa, "0_scatto0.png"), img0)

        m = df.board_rect_mask(img0)
        if m is not None:
            vis = cv2.addWeighted(img0, 0.35, cv2.bitwise_and(img0, img0, mask=m), 0.65, 0)
            cnts, _ = cv2.findContours(m, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
            cv2.drawContours(vis, cnts, -1, ACC, 3)
            self._titolo(vis, "FASE 1 - maschera board (board_rect_mask)")
            cv2.imwrite(os.path.join(self.d_mappa, "1_maschera.png"), vis)
        src = cv2.bitwise_and(img0, img0, mask=m) if m is not None else img0

        model, markers = bm.build_model(src, verbose=False)
        ang = df.marker_angle(markers)
        raddrizzata = abs(ang) > df.DESKEW_MIN
        if raddrizzata:
            M = cv2.getRotationMatrix2D((src.shape[1] / 2, src.shape[0] / 2), ang, 1.0)
            ruotata = cv2.warpAffine(src, M, (src.shape[1], src.shape[0]),
                                     flags=cv2.INTER_CUBIC)
            vis = ruotata.copy()
            self._titolo(vis, f"FASE 3 - deskew {ang:+.2f} gradi (modello RIFATTO)")
            cv2.imwrite(os.path.join(self.d_mappa, "3_deskew.png"), vis)
            src = ruotata
            model, markers = bm.build_model(src, verbose=False)

        pos = og.positions(model)
        vis = src.copy()
        for p in pos.values():
            cv2.circle(vis, p, 3, (0, 200, 0), -1)
        for u, v, x, y in markers[:4]:
            cv2.circle(vis, (int(x), int(y)), 24, (0, 255, 255), 3)
        self._titolo(vis, f"FASE 2 - 4 marker + griglia agganciata ({len(pos)} fori)")
        cv2.imwrite(os.path.join(self.d_mappa, "2_marker_griglia.png"), vis)

        probs = df.map_health(model, markers)
        pitch = abs(pos[(0, 1)][0] - pos[(0, 0)][0]) if (0, 1) in pos and (0, 0) in pos else 0
        self.mappa_info = {"fori": len(pos), "passo_px": pitch,
                           "angolo": round(float(ang), 2), "raddrizzata": raddrizzata,
                           "problemi": list(probs) if probs else []}
        pan = np.full((240, 900, 3), PAPER, np.uint8)
        righe = [f"FASE 4 - salute della mappa",
                 f"fori registrati: {len(pos)}   passo: {pitch}px   angolo: {ang:+.2f} deg"
                 + ("  (RADDRIZZATA)" if raddrizzata else ""),
                 "problemi: " + ("; ".join(probs) if probs else "nessuno - MAPPA OK")]
        for i, t in enumerate(righe):
            cv2.putText(pan, t, (24, 50 + i * 56), FONT, 0.72,
                        INK if i else ACC_D, 2 if not i else 1)
        cv2.imwrite(os.path.join(self.d_mappa, "4_salute.png"), pan)
        json.dump(self.mappa_info, open(os.path.join(self.d_mappa, "mappa.json"),
                                        "w", encoding="utf-8"), indent=2)
        _log(f"mappa: {len(pos)} fori, passo {pitch}px, angolo {ang:+.2f}")

    # ------------------------------------------------------------ 01: scatti
    @_blindato
    def scatto(self, nshot, esito, motivo="", d_prev=None, novelty=None,
               allineamento=None, caldi=None, blob_area=None, img=None, ora=None):
        """Una riga per ogni scatto (analizzati e scartati: i numeri costano nulla).
        Le foto solo per i primi MAX_ESEMPI di ogni motivo di scarto.

        `ora` = quando lo scatto è stato preso. Cattura e analisi sono sfasate: la
        riga di un frame analizzato viene scritta a fine analisi e senza questo
        parametro porterebbe l'ora sbagliata di parecchi secondi."""
        with self._lock:
            self._scatto(nshot, esito, motivo, d_prev, novelty, allineamento,
                         caldi, blob_area, img, ora)

    def _scatto(self, nshot, esito, motivo, d_prev, novelty, allineamento,
                caldi, blob_area, img, ora=None):
        chiave = f"{esito}:{motivo}" if motivo else esito
        self.contatori[chiave] = self.contatori.get(chiave, 0) + 1
        foto = ""
        if img is not None and esito != "ANALIZZATO":
            n = self.n_esempi.get(chiave, 0)
            if n < MAX_ESEMPI:
                self.n_esempi[chiave] = n + 1
                safe = "".join(c if c.isalnum() else "_" for c in chiave)[:40]
                foto = f"{safe}_{nshot:04d}.jpeg"
                vis = img.copy()
                self._titolo(vis, f"#{nshot} {esito} - {motivo}"[:90])
                cv2.imwrite(os.path.join(self.d_scarti, foto), vis)
        with open(self.csv_path, "a", newline="", encoding="utf-8") as f:
            csv.writer(f).writerow(
                [nshot, ora or datetime.datetime.now().strftime("%H:%M:%S"), esito, motivo,
                 _r(d_prev), _r(novelty), _r(allineamento, 1), caldi or "",
                 blob_area or "", foto])

    # ------------------------------------------------------------ 02: valutazioni
    @_blindato
    def valutazione(self, res, rec, ordine, tipo="nuovo"):
        """Artefatti della singola valutazione: blob, viste NCC A+D, crop COSA, record.
        `ordine` = posizione nello storico (1-based): errore e correzione = due voci."""
        import differential as df
        tag = f"v{ordine:03d}"
        step = rec.get("step")
        posx = {**res["pos"], **res["rpos"]}
        thr, weak = res.get("cls_thr") or df.cls_thr(rec.get("atteso", {}).get("cls"))

        # --- blob = solo il diff ---
        p_blob = ""
        if res.get("blob_box") is not None:
            diff = cv2.absdiff(res["img"], res["prev"])
            diff = cv2.convertScaleAbs(diff, alpha=2.2)
            x0, y0, x1, y1 = [int(v) for v in res["blob_box"]]
            try:                             # crop esteso alla sagoma vera (gambe incluse)
                mk = sagoma_estesa(res)
                ys, xs = np.nonzero(mk)
                if len(xs):
                    x0, y0 = min(x0, int(xs.min())), min(y0, int(ys.min()))
                    x1, y1 = max(x1, int(xs.max())), max(y1, int(ys.max()))
            except Exception:
                pass
            g = 80
            H, W = diff.shape[:2]
            zy0, zy1 = max(0, y0 - g), min(H, y1 + g)
            zx0, zx1 = max(0, x0 - g), min(W, x1 + g)
            diff = diff[zy0:zy1, zx0:zx1]
            # solo il diff: la prova fotografica sta già nella vista D
            self._titolo(diff, "diff prev vs ora (x2.2)"[:int(diff.shape[1] / 11)])
            p_blob = f"{tag}_blob.png"
            cv2.imwrite(os.path.join(self.d_val, p_blob), _ridim(diff, 430))

        att_uv = []
        for hname in rec.get("atteso", {}).get("coords", []):
            try:
                att_uv.append(df.parse_hole(hname))
            except Exception:
                pass

        # --- vista A: breadboard sintetica a colonne dritte (colonna 'NCC vista A'
        # dell'Excel) ---
        p_a = f"{tag}_ncc_A.png"
        va = self._vista_a(res, rec, att_uv, posx, thr, weak, df)
        if va is not None:
            cv2.imwrite(os.path.join(self.d_val, p_a), va)
        else:
            p_a = ""

        # --- vista D: foto + overlay misure (prova fotografica) ---
        p_d = f"{tag}_ncc_D.png"
        vd = self._vista_d(res, rec, att_uv, posx, thr, weak, df)
        if vd is not None:
            cv2.imwrite(os.path.join(self.d_val, p_d), vd)
        else:
            p_d = ""

        # --- COSA: crop attorno al blob + chi ha detto la classe ---
        p_cosa = ""
        if res.get("blob_box") is not None:
            x0, y0, x1, y1 = [int(v) for v in res["blob_box"]]
            g = 60
            H, W = res["img"].shape[:2]
            crop = res["img"][max(0, y0 - g):min(H, y1 + g),
                              max(0, x0 - g):min(W, x1 + g)].copy()
            fonte = rec.get("cosa", {}).get("fonte", "?")
            visto = rec.get("cosa", {}).get("cls_vista") or "(non letta)"
            self._titolo(crop, f"COSA: {visto} ({fonte}) - QUESTO e' cio' che il VLM vede")
            p_cosa = f"{tag}_cosa.png"
            cv2.imwrite(os.path.join(self.d_val, p_cosa), crop)

        json.dump({k: v for k, v in rec.items() if k != "report"},
                  open(os.path.join(self.d_val, f"{tag}_rec.json"), "w",
                       encoding="utf-8"), indent=2, ensure_ascii=False)

        self.righe.append({
            "ordine": ordine, "ora": datetime.datetime.now().strftime("%H:%M:%S"),
            "k": res["k"], "tipo": tipo, "step": step,
            "atteso": rec.get("atteso", {}).get("el", ""),
            "fori_attesi": ",".join(rec.get("atteso", {}).get("coords", [])),
            "cls_vista": rec.get("cosa", {}).get("cls_vista") or "",
            "fonte": rec.get("cosa", {}).get("fonte", ""),
            "colore": rec.get("cosa", {}).get("col_vista") or "",
            "verdetto": rec.get("verdetto", ""),
            "pin_promossi": ("-" if (rec.get("atteso") or {}).get("cls") in
                             {"resistenza", "condensatore", "diodo", "led"}
                             else ",".join(rec.get("radar", {}).get("pins", [])) or "-"),
            "spiegazione": (rec.get("spiegazione") or "")[:160],
            "img": {"foto": f"{res['k']}.jpeg", "blob": p_blob,
                    "ncc_A": p_a, "ncc_D": p_d, "cosa": p_cosa},
        })
        try:                                   # crash-safe: append immediato, così una
            with open(os.path.join(self.base, "righe.jsonl"), "a",
                      encoding="utf-8") as f:  # run chiusa a mano non perde le righe
                f.write(json.dumps(self.righe[-1], ensure_ascii=False) + "\n")
        except Exception:
            pass

    # ------------------------------------------------------------ GT live
    @_blindato
    def gt_live(self, ordine, tag):
        """Tasti 1/2/3 durante la run: giudizio dell'operatore sull'ULTIMA valutazione.
        Persistito subito (crash-safe)."""
        self.gt[str(ordine)] = tag
        json.dump(self.gt, open(self._gt_path, "w", encoding="utf-8"), indent=2)
        _log(f"GT live: valutazione {ordine} -> {tag}")

    # ------------------------------------------------------------ chiusura: Excel
    @_blindato
    def chiudi(self, oper=None):
        """RUN_<runid>.xlsx: foglio 'run' (mappa, contatori, pagella) + foglio
        'valutazioni' (una riga per valutazione, miniature incorporate + hyperlink).
        Se le righe in memoria mancano (run uccisa e rigenerazione a posteriori),
        si ricaricano da righe.jsonl."""
        if not self.righe:
            try:
                with open(os.path.join(self.base, "righe.jsonl"), encoding="utf-8") as f:
                    self.righe = [json.loads(r) for r in f if r.strip()]
            except Exception:
                pass
        import openpyxl
        from openpyxl.styles import Font as F, PatternFill, Alignment, Border, Side
        from openpyxl.drawing.image import Image as XImg
        from openpyxl.utils import get_column_letter

        HEAD = PatternFill("solid", fgColor="1A2430")
        HEADF = F(name="Bahnschrift", bold=True, color="FFFFFF", size=10)
        TITLE = F(name="Bahnschrift", bold=True, size=15, color="1A2430")
        SUB = F(name="Consolas", size=8, color="4C5966")
        KO_F = PatternFill("solid", fgColor="F9E3D5")
        OK_F = PatternFill("solid", fgColor="EAF3EC")
        GIA_F = PatternFill("solid", fgColor="FBF3D9")   # giallo: OK?/DEBOLE (semaforo)
        NEU_F = PatternFill("solid", fgColor="F1F3F5")
        thin = Side(style="thin", color="C9D0D6")
        BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
        CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
        LFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

        wb = openpyxl.Workbook()

        # ---------- foglio run ----------
        ws = wb.active
        ws.title = "run"
        ws.cell(1, 1, f"RUN {self.runid}" + (f" - commessa {self.cid}" if self.cid else "")).font = TITLE
        ws.cell(2, 1, "telemetria VLM - una riga per valutazione nel foglio 'valutazioni'; "
                      "fonte di verita': il codice (2-vlm/src)").font = SUB
        r = 4
        ws.cell(r, 1, "MAPPA").font = HEADF; ws.cell(r, 1).fill = HEAD
        for kk, vv in (self.mappa_info or {"mappa": "non registrata"}).items():
            r += 1
            ws.cell(r, 1, kk).border = BORD
            ws.cell(r, 2, str(vv)).border = BORD
        r += 2
        ws.cell(r, 1, "SCATTI (esito:motivo -> conteggio)").font = HEADF
        ws.cell(r, 1).fill = HEAD
        for kk in sorted(self.contatori):
            r += 1
            ws.cell(r, 1, kk).border = BORD
            ws.cell(r, 2, self.contatori[kk]).border = BORD
        r += 2
        ws.cell(r, 1, "PAGELLA (GT live)").font = HEADF; ws.cell(r, 1).fill = HEAD
        gt_vals = list(self.gt.values())
        ko_r = [x for x in self.righe if str(x["verdetto"]).startswith("KO")]
        pagella = [("valutazioni totali", len(self.righe)),
                   ("verdetti KO", len(ko_r)),
                   ("GT live: giusto", gt_vals.count("giusto")),
                   ("GT live: errore vero intercettato", gt_vals.count("errore_vero")),
                   ("GT live: falso allarme", gt_vals.count("falso_allarme")),
                   ("senza GT live", len(self.righe) - len(gt_vals))]
        for kk, vv in pagella:
            r += 1
            ws.cell(r, 1, kk).border = BORD
            ws.cell(r, 2, vv).border = BORD
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 42

        # ---------- foglio valutazioni ----------
        wv = wb.create_sheet("valutazioni")
        H = ["#", "foto", "tipo", "step", "atteso", "fori attesi", "pin promossi",
             "classe vista", "fonte", "verdetto", "controllo operatore",
             "spiegazione", "scatto", "blob", "NCC foto", "NCC vista A"]
        W = [5, 6, 11, 6, 22, 14, 14, 13, 8, 14, 16, 30, 22, 26, 22, 22]
        for i, (h, w) in enumerate(zip(H, W), 1):
            c = wv.cell(1, i, h)
            c.fill = HEAD; c.font = HEADF; c.alignment = CEN; c.border = BORD
            wv.column_dimensions[get_column_letter(i)].width = w
        wv.freeze_panes = "A2"
        GT_TXT = {"giusto": "montaggio giusto", "errore_vero": "ERRORE (beccato)",
                  "falso_allarme": "FALSO ALLARME"}
        for j, x in enumerate(self.righe):
            r = j + 2
            wv.row_dimensions[r].height = 84
            gt = GT_TXT.get(self.gt.get(str(x["ordine"]), ""), "")
            spiega = x["spiegazione"]
            if len(spiega) > 70:                    # in tabella corta; intera nel rec.json
                spiega = spiega[:67] + "..."
            vals = [x["ordine"], x["k"], x["tipo"], x["step"] or "-",
                    x["atteso"], x["fori_attesi"], x.get("pin_promossi", "-"),
                    x["cls_vista"], x["fonte"], x["verdetto"], gt, spiega]
            for i, v in enumerate(vals, 1):
                c = wv.cell(r, i, v)
                c.border = BORD
                c.alignment = LFT if i in (5, 12) else CEN
            # semaforo: stessa regola di colore_verdetto — OK? e DEBOLE sono gialli
            # (controllo a vista), non verdi
            v = str(x["verdetto"])
            fill = (KO_F if v.startswith("KO") else
                    GIA_F if v in ("OK?", "DEBOLE") else
                    OK_F if v.startswith("OK") else NEU_F)
            for i in range(1, len(H) + 1):
                wv.cell(r, i).fill = fill
            # miniature + hyperlink (richiede Pillow)
            imgs = [("foto", os.path.join(self.rundir, x["img"]["foto"])),
                    ("blob", os.path.join(self.d_val, x["img"]["blob"])),
                    ("ncc_D", os.path.join(self.d_val, x["img"]["ncc_D"])),
                    ("ncc_A", os.path.join(self.d_val, x["img"].get("ncc_A") or ""))]
            for col, (nome, path) in enumerate(imgs, 13):
                if not path or not os.path.exists(path) or os.path.isdir(path):
                    continue
                thumb = self._thumb(path, f"{x['ordine']:03d}_{nome}")
                if thumb:
                    im = XImg(thumb)
                    wv.add_image(im, f"{get_column_letter(col)}{r}")
                cell = wv.cell(r, col)
                cell.hyperlink = path
                cell.border = BORD

        out = os.path.join(self.rundir, f"RUN_{self.runid}.xlsx")
        wb.save(out)
        _log(f"Excel scritto: {out}  ({len(self.righe)} valutazioni)")
        return out

    # ------------------------------------------------------------ disegno viste
    def _vista_a(self, res, rec, att_uv, posx, thr, weak, df):
        """Vista A: breadboard sintetica orientata come la foto — righe ordinate per
        posizione reale in pixel (se J sta in alto nella foto, sta in alto anche qui)
        — con i nodi delle rail e i loro NCC. Pieno=caldo, chiaro=debole,
        contorno=muto; golden riquadrato."""
        S = {uv: v for uv, v in (res.get("legscore") or {}).items() if uv in posx}
        for uv in att_uv:
            # Golden sempre con i numeri del verdetto: il foglio mostrava il diff
            # prev mentre il giudizio misurava vs àncora/base, e i due valori non
            # coincidevano. Gli altri fori restano diff prev (contesto).
            mis = rec.get("ancorata", {}).get("misure", {}).get(df.name(*uv))
            if uv in posx and mis is not None:
                S[uv] = float(mis)
            elif uv not in S and uv in posx:
                S[uv] = 0.0
        if not S:
            return None
        # RAIL nella finestra x della regione: nodi con NCC calcolato qui
        try:
            g0 = cv2.cvtColor(res["prev"], cv2.COLOR_BGR2GRAY).astype(np.float32)
            g1 = cv2.cvtColor(res["img"], cv2.COLOR_BGR2GRAY).astype(np.float32)
            xs_reg = [posx[uv][0] for uv in S]
            xlo, xhi = min(xs_reg) - 40, max(xs_reg) + 40
            for uv in res.get("rpos", {}):
                if uv not in S and xlo <= posx[uv][0] <= xhi:
                    S[uv] = float(df.legscore(g0, g1, posx, *uv))
        except Exception:
            pass
        # Layout: righe raggruppate per uv[0] e ordinate per y medio reale (se J sta
        # in alto nella foto, sta in alto anche qui). Colonne al raster logico: la x
        # viene dalla colonna del foro, non dai pixel, così la prospettiva della foto
        # non piega la griglia. Le rail entrano con la colonna di griglia più vicina.
        righe_k = {}
        for uv in S:
            righe_k.setdefault(uv[0], []).append(uv)
        ordine_righe = sorted(righe_k, key=lambda kk: np.mean([posx[q][1]
                                                               for q in righe_k[kk]]))
        riga_y = {kk: i for i, kk in enumerate(ordine_righe)}
        vs = sorted({uv[1] for uv in S})
        v_min = vs[0]
        cell, mx, my, leg = 46, 64, 56, 210
        n_col = max(1, vs[-1] - v_min)
        W = mx * 2 + n_col * cell + leg
        Hh = my * 2 + (len(ordine_righe) - 1) * cell + 40
        can = np.full((Hh, W, 3), PAPER, np.uint8)

        def pxy(uv):
            return (mx + (uv[1] - v_min) * cell, my + riga_y[uv[0]] * cell)
        # etichette riga (lettera o segno rail) a sinistra; colonne dai nomi griglia
        for kk in ordine_righe:
            lab = kk[0] if isinstance(kk, str) else df.name(kk, 0)[:1]
            cv2.putText(can, lab, (mx - 44, my + riga_y[kk] * cell + 6),
                        FONT, 0.55, INK2 if isinstance(kk, int) else ACC_D, 2)
        vis_cols = sorted({uv[1] for uv in S if isinstance(uv[0], int)})
        for v in vis_cols[::max(1, len(vis_cols) // 8)]:
            uvq = next((q for q in S if isinstance(q[0], int) and q[1] == v), None)
            if uvq:
                cv2.putText(can, df.name(*uvq)[1:], (pxy(uvq)[0] - 8, my - 26),
                            FONT, 0.45, INK2, 1)
        rim = _rimozione_uv(rec, posx, df, att_uv)
        for uv, val in S.items():
            x, y = pxy(uv)
            if uv in rim and val >= weak:
                # zona del KO precedente: caldo di RIMOZIONE, non un pezzo -> grigio
                cv2.circle(can, (x, y), 9, HAIR, -1)
                cv2.circle(can, (x, y), 9, INK2, 1)
            elif val >= thr:
                cv2.circle(can, (x, y), 9, ACC, -1)
                cv2.circle(can, (x, y), 9, ACC_D, 2)
            elif val >= weak:
                cv2.circle(can, (x, y), 9, ACC_L, -1)
                cv2.circle(can, (x, y), 9, ACC, 1)
            else:
                cv2.circle(can, (x, y), 9, HAIR, 2)
            if val >= weak or uv in att_uv:
                col_t = INK2 if uv in rim else INK
                cv2.putText(can, f"{val:.2f}", (x - 18, y - 15), FONT, 0.42, col_t, 1)
        for uv in att_uv:
            if uv in S:
                x, y = pxy(uv)
                cv2.rectangle(can, (x - 15, y - 15), (x + 15, y + 15), INK, 2)
        # Niente X per le classi in cui il pick libero non decide il verdetto:
        # solo-golden, corpo, filo (il pick aggancia la cupola del led vicino o la
        # rail attraversata: decide il ramo filo) e transistor (decide il ramo
        # percorso). Lì la X sarebbe solo rumore.
        no_x = (rec.get("atteso") or {}).get("cls") in {
            "resistenza", "condensatore", "diodo", "led", "transistor",
            "trimmer", "buzzer", "potenziometro", "jumper", "cavetto"}
        fant = _fantasma_uv(rec, df)
        for uv in ([] if no_x else (res.get("pins") or [])):   # X solo dove il radar decide
            if isinstance(uv[0], int) and uv in S:
                x, y = pxy(uv)
                # fantasma del KO precedente: il pick l'ha agganciato ma il giudizio
                # lo ignora (rimozione, non pezzo) -> X grigia, non "scelto"
                col_x = HAIR if uv in fant else ACC_D
                for dx, dy in ((-1, -1), (1, -1)):
                    cv2.line(can, (x - 7 * dx, y - 7 * dy), (x + 7 * dx, y + 7 * dy),
                             col_x, 2)
        # legenda
        lx = W - leg + 18
        voci = [("caldo >= %.2f" % thr, ACC, -1), ("debole >= %.2f" % weak, ACC_L, -1),
                ("muto", HAIR, 2)]
        for i, (txt, col, fl) in enumerate(voci):
            cv2.circle(can, (lx, 40 + i * 30), 8, col, fl)
            if fl == -1:
                cv2.circle(can, (lx, 40 + i * 30), 8, ACC_D if col == ACC else ACC, 1)
            cv2.putText(can, txt, (lx + 18, 45 + i * 30), FONT, 0.42, INK2, 1)
        cv2.rectangle(can, (lx - 8, 122), (lx + 8, 138), INK, 2)
        cv2.putText(can, "golden atteso", (lx + 18, 135), FONT, 0.42, INK2, 1)
        if not no_x:                       # niente voce X per le classi solo-golden
            for dx, dy in ((-1, -1), (1, -1)):
                cv2.line(can, (lx - 6 * dx, 158 - 6 * dy), (lx + 6 * dx, 158 + 6 * dy),
                         ACC_D, 2)
            cv2.putText(can, "pin scelti dal sistema", (lx + 18, 162), FONT, 0.42, INK2, 1)
            if fant:
                for dx, dy in ((-1, -1), (1, -1)):
                    cv2.line(can, (lx - 6 * dx, 188 - 6 * dy), (lx + 6 * dx, 188 + 6 * dy),
                             HAIR, 2)
                cv2.putText(can, "fantasma del KO (ignorati)", (lx + 18, 192),
                            FONT, 0.42, INK2, 1)
        if rim:
            cv2.circle(can, (lx, 218), 8, HAIR, -1)
            cv2.circle(can, (lx, 218), 8, INK2, 1)
            cv2.putText(can, "rimozione del KO prec.", (lx + 18, 222),
                        FONT, 0.42, INK2, 1)
        el = rec.get("atteso", {}).get("el", "")
        cv2.putText(can, f"{el}"[:46], (18, Hh - 14), FONT, 0.5, INK, 1)
        return can

    def _vista_d(self, res, rec, att_uv, posx, thr, weak, df):
        """Vista D: tre strati leggibili in ordine — 1) fori caldi (cerchio + valore),
        2) golden atteso (quadrato), 3) pin promossi da fit_pins (X). Valori solo su
        caldi/golden/promossi, per non sovrapporre troppo."""
        img = res.get("img")
        if img is None:
            return None
        S = {uv: v for uv, v in (res.get("legscore") or {}).items()
             if isinstance(uv[0], int) and uv in posx}
        # Niente X per le classi in cui il pick libero non decide il verdetto
        # (solo-golden, corpo, filo, transistor): v. _vista_a.
        SOLO_GOLDEN = {"resistenza", "condensatore", "diodo", "led", "transistor",
                       "trimmer", "buzzer", "potenziometro", "jumper", "cavetto"}
        no_x = (rec.get("atteso") or {}).get("cls") in SOLO_GOLDEN
        pins = [] if no_x else [uv for uv in (res.get("pins") or [])
                                if isinstance(uv[0], int) and uv in posx]
        rail_p = [uv for uv in (res.get("pins") or [])
                  if not isinstance(uv[0], int) and uv in posx]
        rilevanti = ([uv for uv, v in S.items() if v >= weak]
                     + [uv for uv in att_uv if uv in posx] + pins + rail_p)
        if not rilevanti:
            return None
        xs = [posx[uv][0] for uv in rilevanti]; ys = [posx[uv][1] for uv in rilevanti]
        pitch = res.get("pitch") or 30
        g = int(4 * pitch)
        H, W = img.shape[:2]
        x0, x1 = max(0, int(min(xs)) - g), min(W, int(max(xs)) + g)
        y0, y1 = max(0, int(min(ys)) - g), min(H, int(max(ys)) + g)
        vis = img[y0:y1, x0:x1].copy()

        def put(txt, x, y, col, sc=0.5):
            cv2.putText(vis, txt, (x, y), FONT, sc, (20, 25, 30), 4)
            cv2.putText(vis, txt, (x, y), FONT, sc, col, 1)

        def pj(uv):
            return int(posx[uv][0]) - x0, int(posx[uv][1]) - y0
        # strato 1: misure — al massimo 6 fori caldi, per leggibilità
        caldi6 = sorted((uv for uv, v in S.items() if v >= weak),
                        key=lambda q: -S[q])[:6]
        rim = _rimozione_uv(rec, posx, df, att_uv)
        for uv in set(caldi6) | (set(att_uv) & set(S)) | set(pins):
            v = S.get(uv, 0.0)
            x, y = pj(uv)
            if uv in rim and v >= weak:
                # caldo di RIMOZIONE (zona del KO precedente): grigio, non arancio
                cv2.circle(vis, (x, y), 13, (160, 160, 160), 2)
                put(f"{v:.2f}", x - 18, y - 18, (170, 170, 170), 0.48)
                continue
            if v >= thr:
                cv2.circle(vis, (x, y), 13, ACC, 3)
            elif v >= weak:
                cv2.circle(vis, (x, y), 11, (95, 162, 232), 2)
            put(f"{v:.2f}", x - 18, y - 18, (255, 255, 255), 0.48)
        # capo su rail: i pin non-griglia del pick si marcano con la loro polarità
        for uv in ([] if no_x else (res.get("pins") or [])):
            if not isinstance(uv[0], int) and uv in posx:
                x, y = pj(uv)
                if -40 < x < vis.shape[1] + 40 and -40 < y < vis.shape[0] + 40:
                    for dx, dy in ((-1, -1), (1, -1)):
                        cv2.line(vis, (x - 9 * dx, y - 9 * dy), (x + 9 * dx, y + 9 * dy),
                                 (60, 230, 255), 3)
                    put(f"rail {uv[0][0]}", x - 24, y + 30, (60, 230, 255), 0.55)
        for uv in att_uv:                                # strato 2: golden
            if uv in posx:
                x, y = pj(uv)
                cv2.rectangle(vis, (x - 17, y - 17), (x + 17, y + 17), (255, 255, 255), 2)
        fant = _fantasma_uv(rec, df)
        for uv in pins:                                  # strato 3: pin promossi (X)
            x, y = pj(uv)
            col_x = (150, 150, 150) if uv in fant else (60, 230, 255)
            for dx, dy in ((-1, -1), (1, -1)):
                cv2.line(vis, (x - 9 * dx, y - 9 * dy), (x + 9 * dx, y + 9 * dy),
                         (20, 25, 30), 5)
                cv2.line(vis, (x - 9 * dx, y - 9 * dy), (x + 9 * dx, y + 9 * dy),
                         col_x, 2)
        t = (f"{rec.get('verdetto', '')} | O caldo  [] golden"
             + ("" if no_x else "  X scelti"))
        if fant or rim:
            t += "  (grigio = rimozione/fantasma del KO prec.)"
        cv2.rectangle(vis, (0, 0), (vis.shape[1], 30), (30, 25, 20), -1)
        cv2.putText(vis, t[:int(vis.shape[1] / 11)], (8, 21), FONT, 0.5, (240, 237, 232), 1)
        return vis

    # ------------------------------------------------------------ util
    def _titolo(self, img, txt):
        cv2.rectangle(img, (0, 0), (img.shape[1], 34), (30, 25, 20), -1)
        cv2.putText(img, txt, (10, 24), FONT, 0.6, (240, 237, 232), 1)

    def _thumb(self, path, nome):
        """Miniatura PNG per l'Excel (cartella _thumbs, larghezza THUMB_W)."""
        try:
            im = cv2.imread(path)
            if im is None:
                return None
            s = THUMB_W / im.shape[1]
            im = cv2.resize(im, (THUMB_W, max(1, int(im.shape[0] * s))),
                            interpolation=cv2.INTER_AREA)
            d = os.path.join(self.base, "_thumbs")
            os.makedirs(d, exist_ok=True)
            p = os.path.join(d, f"{nome}.png")
            cv2.imwrite(p, im)
            return p
        except Exception:
            return None


def sagoma_estesa(res, soglia=18):
    """Delega a filo.sagoma_diff: una sola implementazione, la telemetria la eredita."""
    import filo as fl
    return fl.sagoma_diff(res, soglia)


def _spina(mask, c2, sc=4):
    """La spina del pezzo: cammino capo->capo dentro la sagoma, pesato sulla
    distanza dal bordo (Dijkstra su griglia sottocampionata), così la linea corre
    al centro del metallo, non lungo il contorno."""
    import heapq
    small = cv2.resize(mask, None, fx=1.0 / sc, fy=1.0 / sc,
                       interpolation=cv2.INTER_NEAREST)
    dt = cv2.distanceTransform((small > 0).astype(np.uint8), cv2.DIST_L2, 3)
    a = (c2[0][0] // sc, c2[0][1] // sc)
    b = (c2[1][0] // sc, c2[1][1] // sc)
    h, w = small.shape
    mx = float(dt.max()) or 1.0
    dist = {a: 0.0}
    prev = {}
    q = [(0.0, a)]
    while q:
        d, p = heapq.heappop(q)
        if p == b:
            break
        if d > dist.get(p, 1e18):
            continue
        x, y = p
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                if not dx and not dy:
                    continue
                nx, ny = x + dx, y + dy
                if not (0 <= nx < w and 0 <= ny < h) or not small[ny, nx]:
                    continue
                nd = d + 0.05 + (mx - dt[ny, nx]) / mx     # centro = costo minimo
                if nd < dist.get((nx, ny), 1e18):
                    dist[(nx, ny)] = nd
                    prev[(nx, ny)] = p
                    heapq.heappush(q, (nd, (nx, ny)))
    if b not in dist:
        return None
    path, p = [], b
    while p != a:
        path.append((p[0] * sc, p[1] * sc))
        p = prev[p]
    path.append((a[0] * sc, a[1] * sc))
    return path[::-1]


def _r(v, nd=4):
    return round(float(v), nd) if v is not None else ""


def _ridim(img, h):
    s = h / img.shape[0]
    return cv2.resize(img, (int(img.shape[1] * s), h))
