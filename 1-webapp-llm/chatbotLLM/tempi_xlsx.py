"""tempi_xlsx.py — mette i tempi di fase di una commessa in data/tempi_commesse.xlsx.

I tempi li stampa il server (POST /api/tempi) nel terminale, in due righe:
    [tempi] commessa D5  |  T1 picking 01:12  T2 assemblaggio 09:40  T3 collaudo 02:05  ...
    [tempi] D5	72.4	580.3	125.1	777.8
Questo script prende i secondi e li scrive nel foglio, una riga per commessa.
Rilanciabile: se la commessa c'è già, la riga viene aggiornata invece che duplicata.

Uso:
    python tempi_xlsx.py D5 72.4 580.3 125.1
    python tempi_xlsx.py --tsv "D5	72.4	580.3	125.1"     # riga incollata dal terminale
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "tempi_commesse.xlsx")
COLONNE = ["Commessa", "T1 picking (s)", "T2 assemblaggio (s)", "T3 collaudo (s)",
           "Totale (s)", "T1 mm:ss", "T2 mm:ss", "T3 mm:ss", "Totale mm:ss"]


def mmss(sec):
    sec = int(round(max(0.0, sec)))
    return f"{sec // 60:02d}:{sec % 60:02d}"


def foglio():
    """Apre il workbook o lo crea con l'intestazione."""
    if os.path.exists(XLSX):
        wb = openpyxl.load_workbook(XLSX)
        return wb, wb.active
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "tempi"
    ws.cell(1, 1, "TEMPI DI FASE PER COMMESSA").font = Font(name="Bahnschrift", bold=True, size=14)
    ws.cell(2, 1, "cronometro della webapp: T1 picking, T2 assemblaggio, T3 collaudo & CQ").font = \
        Font(name="Consolas", size=8, color="4C5966")
    head = PatternFill("solid", fgColor="1A2430")
    for i, nome in enumerate(COLONNE, 1):
        c = ws.cell(4, i, nome)
        c.font = Font(name="Bahnschrift", bold=True, color="FFFFFF", size=10)
        c.fill = head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 17
    ws.freeze_panes = "A5"
    return wb, ws


def scrivi(cid, t1, t2, t3):
    wb, ws = foglio()
    tot = t1 + t2 + t3
    riga = None
    for r in range(5, ws.max_row + 1):                 # commessa già presente -> aggiorno
        if str(ws.cell(r, 1).value or "").strip().upper() == cid.upper():
            riga = r
            break
    if riga is None:
        riga = max(5, ws.max_row + 1)
    thin = Side(style="thin", color="C9D0D6")
    valori = [cid.upper(), round(t1, 1), round(t2, 1), round(t3, 1), round(tot, 1),
              mmss(t1), mmss(t2), mmss(t3), mmss(tot)]
    for i, v in enumerate(valori, 1):
        c = ws.cell(riga, i, v)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        c.alignment = Alignment(horizontal="center", vertical="center")
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    print(f"{cid.upper()}: T1 {mmss(t1)}  T2 {mmss(t2)}  T3 {mmss(t3)}  totale {mmss(tot)}"
          f"  -> riga {riga} di {XLSX}")


def main():
    a = sys.argv[1:]
    if a and a[0] == "--tsv":
        campi = a[1].replace("[tempi]", "").strip().split("\t")
        if len(campi) < 4:
            sys.exit("riga TSV incompleta: serve commessa + 3 tempi")
        cid, t1, t2, t3 = campi[0].strip(), *[float(x) for x in campi[1:4]]
    elif len(a) == 4:
        cid, t1, t2, t3 = a[0], float(a[1]), float(a[2]), float(a[3])
    else:
        sys.exit(__doc__)
    scrivi(cid, t1, t2, t3)


if __name__ == "__main__":
    main()
